"""
Adedoyin Mobility and Logistics - Operational Data Generator
Student Name: Adedoyin Ajeyomi
Student ID: 20260520
Course: Advanced Database Systems

This script creates realistic operational records for Adedoyin Mobility and Logistics.
The company originally stored all daily trips, bookings, and parcel records in Excel sheets.
We generate that legacy spreadsheet data here so we can load it into our database and test query performance.
"""

import os
import random
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from faker import Faker

fake = Faker(['en_US'])
random.seed(42)

# Set up paths for output files
base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
data_dir = os.path.join(base_dir, 'data')
os.makedirs(data_dir, exist_ok=True)

excel_path = os.path.join(data_dir, 'Ajeyomi_Adedoyin_20260520_Adedoyin_Mobility_Data.xlsx')

wb = openpyxl.Workbook()
wb.remove(wb.active)

# Simple table header styling
header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

def style_sheet(ws, headers):
    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

def adjust_columns(ws):
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

# 1. Stations (6 regional hubs across Nigeria)
print("Generating stations data...")
ws_stations = wb.create_sheet(title="stations")
station_headers = ["station_id", "station_name", "state", "address", "contact_phone", "manager_name"]
style_sheet(ws_stations, station_headers)

stations_data = [
    (1, "Lagos Ojota Terminal", "Lagos", "12 Ikorodu Road, Ojota, Lagos", "08021110001", "Babatunde Adeleke"),
    (2, "Ibadan Iwo Road Terminal", "Oyo", "45 Iwo Road Interchange, Ibadan", "08021110002", "Folashade Alao"),
    (3, "Abuja Utako Central Hub", "FCT", "8 Obafemi Awolowo Way, Utako, Abuja", "08021110003", "Musa Garba"),
    (4, "Benin Ramat Park Terminal", "Edo", "19 Agbor Road, Ramat Park, Benin City", "08021110004", "Osaro Igbinosa"),
    (5, "Enugu Holy Ghost Terminal", "Enugu", "3 Ogui Road, Holy Ghost, Enugu", "08021110005", "Chinedu Okoye"),
    (6, "Kano Sabon Gari Terminal", "Kano", "52 France Road, Sabon Gari, Kano", "08021110006", "Ibrahim Danladi")
]

for row in stations_data:
    ws_stations.append(list(row))
adjust_columns(ws_stations)

# 2. Vehicles (20 fleet vehicles: HiAce buses and Sienna minivans)
print("Generating fleet vehicles...")
ws_vehicles = wb.create_sheet(title="vehicles")
vehicle_headers = ["vehicle_id", "registration_number", "chassis_number", "make_model", 
                   "seating_capacity", "current_station_id", "roadworthiness_expiry", "gps_tracker_id", "status"]
style_sheet(ws_vehicles, vehicle_headers)

vehicles_data = []
nigerian_state_codes = ["KJA", "APP", "ABC", "RBC", "ENU", "KAN", "IBD", "BEN"]

for v_id in range(1, 21):
    plate = f"{random.choice(nigerian_state_codes)}-{random.randint(100, 999)}{random.choice(['AA', 'XY', 'KT', 'LG', 'AB'])}"
    vin = f"2T1BR32E{random.randint(100000000, 999999999)}"
    if v_id <= 12:
        model = "Toyota HiAce Bus"
        cap = random.choice([14, 16, 18])
    else:
        model = "Toyota Sienna Minivan"
        cap = random.choice([7, 8])
    station = random.randint(1, 6)
    expiry = (datetime(2026, 10, 1) + timedelta(days=random.randint(60, 365))).strftime("%Y-%m-%d")
    tracker = f"GPS-AML-{v_id:03d}"
    status = "Active" if v_id <= 18 else "Under-Maintenance"
    
    vehicles_data.append((v_id, plate, vin, model, cap, station, expiry, tracker, status))
    ws_vehicles.append(list(vehicles_data[-1]))

adjust_columns(ws_vehicles)

# 3. Drivers (20 drivers with verified NIN and licenses)
print("Generating drivers...")
ws_drivers = wb.create_sheet(title="drivers")
driver_headers = ["driver_id", "full_name", "phone_number", "nin", "license_number", 
                  "license_expiry", "assigned_vehicle_id", "status"]
style_sheet(ws_drivers, driver_headers)

first_names = ["Kehinde", "Emeka", "Suleiman", "Oluwaseun", "Sunday", "Aliyu", "Godwin", "Kayode", 
               "Tunde", "Ifeanyi", "Abubakar", "Samuel", "Victor", "Femi", "Chidi", "Haruna", "Jamiu", "David", "Chukwuma", "Bashir"]
last_names = ["Ogunleye", "Okafor", "Bello", "Ajayi", "Nnamani", "Shehu", "Adeyemi", "Eze", 
              "Bakare", "Danjuma", "Lawal", "Nwosu", "Salami", "Ojo", "Idris", "Afolabi", "Balogun", "Kalu", "Gbadamosi", "Yakubu"]

drivers_data = []
for d_id in range(1, 21):
    name = f"{first_names[d_id-1]} {last_names[d_id-1]}"
    phone = f"080{random.choice(['3', '5', '7', '8'])}{random.randint(1000000, 9999999)}"
    nin = f"{random.randint(10000000000, 99999999999)}"
    license_no = f"FRSC-{d_id:04d}-{random.choice(['LA', 'ABJ', 'OY', 'KN', 'EN'])}"
    expiry = (datetime(2026, 10, 1) + timedelta(days=random.randint(90, 700))).strftime("%Y-%m-%d")
    assigned_veh = d_id
    status = "Available" if d_id <= 18 else "On-Leave"
    
    drivers_data.append((d_id, name, phone, nin, license_no, expiry, assigned_veh, status))
    ws_drivers.append(list(drivers_data[-1]))

adjust_columns(ws_drivers)

# 4. Customers (500 customers with contact and emergency details)
print("Generating customer profiles...")
ws_customers = wb.create_sheet(title="customers")
customer_headers = ["customer_id", "full_name", "phone_number", "email", "nin", 
                    "emergency_contact_name", "emergency_contact_phone"]
style_sheet(ws_customers, customer_headers)

customers_data = []
for c_id in range(1, 501):
    fname = random.choice(first_names)
    lname = random.choice(last_names)
    full_name = f"{fname} {lname}"
    phone = f"081{random.randint(10000000, 99999999)}"
    email = f"{fname.lower()}.{lname.lower()}{c_id}@gmail.com"
    nin = f"{random.randint(10000000000, 99999999999)}"
    nok_name = f"{random.choice(first_names)} {lname}"
    nok_phone = f"070{random.randint(10000000, 99999999)}"
    
    customers_data.append((c_id, full_name, phone, email, nin, nok_name, nok_phone))
    ws_customers.append(list(customers_data[-1]))

adjust_columns(ws_customers)

# 5. Routes (10 interstate travel corridors)
print("Generating routes...")
ws_routes = wb.create_sheet(title="routes")
route_headers = ["route_id", "origin_station_id", "destination_station_id", 
                 "distance_km", "est_travel_hours", "base_passenger_fare", "base_parcel_rate"]
style_sheet(ws_routes, route_headers)

routes_raw = [
    (1, 1, 2, 130.0, 2.5, 3500.0, 1500.0),    # Lagos to Ibadan
    (2, 2, 1, 130.0, 2.5, 3500.0, 1500.0),    # Ibadan to Lagos
    (3, 1, 3, 750.0, 10.0, 18500.0, 4500.0),  # Lagos to Abuja
    (4, 3, 1, 750.0, 10.0, 18500.0, 4500.0),  # Abuja to Lagos
    (5, 1, 4, 320.0, 5.5, 9000.0, 2500.0),    # Lagos to Benin City
    (6, 4, 1, 320.0, 5.5, 9000.0, 2500.0),    # Benin City to Lagos
    (7, 4, 5, 230.0, 4.0, 6500.0, 2000.0),    # Benin City to Enugu
    (8, 5, 4, 230.0, 4.0, 6500.0, 2000.0),    # Enugu to Benin City
    (9, 3, 6, 440.0, 6.0, 11000.0, 3000.0),   # Abuja to Kano
    (10, 6, 3, 440.0, 6.0, 11000.0, 3000.0)   # Kano to Abuja
]

for r in routes_raw:
    ws_routes.append(list(r))
adjust_columns(ws_routes)

# 6. Trips (1,200 scheduled journeys across routes)
print("Generating 1,200 trips...")
ws_trips = wb.create_sheet(title="trips")
trip_headers = ["trip_id", "route_id", "vehicle_id", "driver_id", 
                "departure_datetime", "arrival_datetime", "trip_status"]
style_sheet(ws_trips, trip_headers)

trips_data = []
start_date = datetime(2026, 6, 1, 6, 0, 0)

for t_id in range(1, 1201):
    route = random.choice(routes_raw)
    route_id = route[0]
    est_hours = route[4]
    
    veh_id = random.randint(1, 20)
    driver_id = random.randint(1, 20)
    
    dep_time = start_date + timedelta(days=t_id // 10, hours=random.choice([6, 7, 9, 12, 14, 20]), minutes=random.choice([0, 15, 30]))
    arr_time = dep_time + timedelta(hours=est_hours + random.uniform(0.2, 1.5))
    
    if dep_time > datetime(2026, 10, 25):
        status = "Scheduled"
    elif dep_time > datetime(2026, 10, 24):
        status = "In-Transit"
    else:
        status = random.choice(["Completed", "Completed", "Completed", "Cancelled"])
        
    trips_data.append((t_id, route_id, veh_id, driver_id, dep_time.strftime("%Y-%m-%d %H:%M:%S"), arr_time.strftime("%Y-%m-%d %H:%M:%S"), status))
    ws_trips.append(list(trips_data[-1]))

adjust_columns(ws_trips)

# 7. Bookings (8,000 passenger seat bookings)
print("Generating 8,000 bookings...")
ws_bookings = wb.create_sheet(title="bookings")
booking_headers = ["booking_id", "trip_id", "customer_id", "seat_number", 
                   "ticket_amount", "booking_status", "payment_reference", "booking_timestamp"]
style_sheet(ws_bookings, booking_headers)

b_id = 1
for trip in trips_data:
    t_id = trip[0]
    route_id = trip[1]
    veh_id = trip[2]
    dep_time = datetime.strptime(trip[4], "%Y-%m-%d %H:%M:%S")
    trip_status = trip[6]
    
    veh_cap = vehicles_data[veh_id - 1][4]
    route_fare = routes_raw[route_id - 1][5]
    
    num_passengers = random.randint(min(5, veh_cap), veh_cap)
    
    for seat in range(1, num_passengers + 1):
        cust_id = random.randint(1, 500)
        pay_ref = f"PAY-AML-{b_id:06d}-{random.randint(100, 999)}"
        b_time = dep_time - timedelta(days=random.randint(1, 5), hours=random.randint(1, 10))
        b_status = "Cancelled" if trip_status == "Cancelled" else "Confirmed"
        
        ws_bookings.append([b_id, t_id, cust_id, seat, route_fare, b_status, pay_ref, b_time.strftime("%Y-%m-%d %H:%M:%S")])
        b_id += 1
        if b_id > 8000:
            break
    if b_id > 8000:
        break

adjust_columns(ws_bookings)

# 8. Parcels (3,500 parcel waybill deliveries)
print("Generating 3,500 parcel deliveries...")
ws_parcels = wb.create_sheet(title="parcels")
parcel_headers = ["parcel_id", "trip_id", "sender_customer_id", "recipient_name", 
                  "recipient_phone", "recipient_nin", "waybill_number", "weight_kg", "declared_value", "delivery_status"]
style_sheet(ws_parcels, parcel_headers)

p_id = 1
for trip in trips_data:
    t_id = trip[0]
    trip_status = trip[6]
    
    for _ in range(random.randint(2, 4)):
        sender = random.randint(1, 500)
        rec_name = f"{random.choice(first_names)} {random.choice(last_names)}"
        rec_phone = f"080{random.randint(10000000, 99999999)}"
        rec_nin = f"{random.randint(10000000000, 99999999999)}"
        waybill = f"AML-WB-2026-{p_id:05d}"
        weight = round(random.uniform(1.5, 45.0), 2)
        val = round(random.choice([15000, 35000, 80000, 150000, 250000, 500000]), 2)
        
        if trip_status == "Completed":
            del_status = random.choice(["Delivered", "Delivered", "Arrived-Hub"])
        elif trip_status == "In-Transit":
            del_status = "In-Transit"
        elif trip_status == "Cancelled":
            del_status = "Returned"
        else:
            del_status = "Registered"
            
        ws_parcels.append([p_id, t_id, sender, rec_name, rec_phone, rec_nin, waybill, weight, val, del_status])
        p_id += 1
        if p_id > 3500:
            break
    if p_id > 3500:
        break

adjust_columns(ws_parcels)

# Save master Excel workbook
print(f"Saving workbook to {excel_path}...")
wb.save(excel_path)

# Export individual CSV files for direct SQL loading
print("Exporting individual CSV files...")
import csv

for sheetname in wb.sheetnames:
    sheet = wb[sheetname]
    csv_file = os.path.join(data_dir, f"{sheetname}.csv")
    with open(csv_file, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        for row in sheet.iter_rows(values_only=True):
            writer.writerow(row)

print("Done! All datasets saved successfully.")
